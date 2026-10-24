"""Case objectives and consulting-framework views of the impact (used by build_model.py).

The case objectives (profitability, growth, environmental, social) are the headline: each KPI is shown for the
base case (do nothing) and with the strategies, and the change is split by pillar and strategy. The same
impact is then cut through the standard consulting frameworks (triple bottom line, profitability tree, Ansoff,
marketing mix, value chain, balanced scorecard), so every view adds up to the same total. NPV stays in the
model as the valuation, not the headline.
"""
from __future__ import annotations

from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter as L

OBJ_KINDS = ("profitability", "growth", "environmental", "social", "other")
QUANTITIES = {"volume": "volume", "revenue": "incremental revenue", "saving": "cost saving", "opex": "opex",
              "capex": "capex", "ebit": "incremental EBIT"}
TBL = {"profitability": "Profit", "growth": "Profit", "environmental": "Planet", "social": "People",
       "other": "Other"}
ANSOFF = {"penetration": "Existing products, existing markets (selling more into today's market)",
          "product_development": "New products, existing markets (product development)",
          "market_development": "Existing products, new markets (market development)",
          "diversification": "New products, new markets (diversification)",
          "efficiency": "Efficiency (cost and resources)"}
FOUR_P = {"price": "Price", "product": "Product", "place": "Place (channel)", "promotion": "Promotion",
          "none": "Not a marketing-mix lever (cost, operations)"}
VALUE_CHAIN = {"inbound_logistics": "Inbound logistics", "operations": "Operations",
               "outbound_logistics": "Outbound logistics", "marketing_sales": "Marketing & sales",
               "service": "Service", "procurement": "Procurement", "technology": "Technology development",
               "hr": "Human resource management", "infrastructure": "Firm infrastructure"}
LENS_TAGS = {"ansoff": ANSOFF, "four_p": FOUR_P, "value_chain": VALUE_CHAIN}
NOT_TAGGED = "Not tagged (set frameworks.<lens> on the initiative)"
CUR_HINTS = ("EUR", "USD", "GBP", "RON", "CHF", "PLN", "€", "$", "£")


class SpecError(Exception):
    pass


def leaf_kind(lf, cur):
    """A price leaf (money per unit: ADR, fee, spend per visit) or a volume leaf (customers, visits, rooms)."""
    if lf.get("kind") in ("price", "volume"):
        return lf["kind"]
    u = str(lf.get("unit", "")).upper()
    return "price" if any(h in u for h in CUR_HINTS + (cur[:3].upper(),)) else "volume"


def normalise_objectives(s, by, derived, inits, leaves, warn):
    """Validate spec objectives and initiative framework tags; fill defaults. Returns input ids used."""
    used = set()
    cur = s["meta"]["currency"]
    objs = s.setdefault("objectives", [])
    if not objs:
        objs.append({"id": "profit", "name": "Profitability", "kind": "profitability"})
        warn.append("no objectives in the spec - headline defaults to profitability; add the case's objectives "
                    "(profitability, growth, environmental, social) from the brief")
    seen = set()
    iby = {it["id"]: it for it in inits}
    for o in objs:
        if o.get("kind") not in OBJ_KINDS:
            raise SpecError(f"objective {o.get('id')}: kind must be one of {', '.join(OBJ_KINDS)}")
        o.setdefault("name", o["id"])
        if o["kind"] in ("profitability", "growth"):
            if o["kind"] == "profitability" and not any(a.get("role") == "ebitda" for a in s.get("actuals", [])):
                warn.append("profitability objective without reported EBITDA (actuals role 'ebitda') - the headline "
                            "shows the profit uplift only, not EBITDA and margin against the base case")
            continue   # built-in KPIs (EBITDA and margin / revenue)
        if not o.get("kpis"):
            raise SpecError(f"objective {o['id']}: a {o['kind']} objective needs kpis "
                            f"[{{id, name, unit, direction, baseline, effects}}]")
        for k in o["kpis"]:
            if not k.get("id") or k["id"] in seen:
                raise SpecError(f"objective {o['id']}: every KPI needs a unique id")
            seen.add(k["id"])
            if k.get("direction") not in ("higher", "lower"):
                raise SpecError(f"KPI {k['id']}: direction must be 'higher' or 'lower' (which way is better)")
            for x in (k.get("baseline") or []) + ([k["baseline_growth"]] if k.get("baseline_growth") else []):
                if x not in by and x not in derived:
                    raise SpecError(f"KPI {k['id']}: baseline input '{x}' is not an input or derived id")
                used.add(x)
            if not k.get("effects"):
                raise SpecError(f"KPI {k['id']}: needs effects [{{initiative, quantity, factor, per, sign}}]")
            for e in k["effects"]:
                it = iby.get(e.get("initiative"))
                if not it:
                    raise SpecError(f"KPI {k['id']}: effect initiative '{e.get('initiative')}' is unknown")
                q = e.get("quantity")
                if q not in QUANTITIES:
                    raise SpecError(f"KPI {k['id']}: quantity must be one of {', '.join(QUANTITIES)}")
                if q == "volume" and it.get("model") != "build":
                    raise SpecError(f"KPI {k['id']}: quantity 'volume' needs a build initiative ({it['id']} is not)")
                if q == "saving" and not (it.get("model") == "build" and it.get("kind") == "saving"):
                    raise SpecError(f"KPI {k['id']}: quantity 'saving' needs a saving build ({it['id']} is not)")
                if q in ("opex", "capex") and not it.get(q):
                    raise SpecError(f"KPI {k['id']}: initiative {it['id']} has no {q}")
                if e.get("sign", 1) not in (1, -1):
                    raise SpecError(f"KPI {k['id']}: sign must be 1 or -1")
                for x in (e.get("factor") or []) + (e.get("per") or []):
                    if x not in by and x not in derived:
                        raise SpecError(f"KPI {k['id']}: factor '{x}' is not an input or derived id (source it)")
                    used.add(x)
    # framework tags: inferred where the model structure says it, otherwise the team sets them
    untagged = {}
    for it in inits:
        fw = it.setdefault("frameworks", {})
        lf = leaves.get(it.get("leaf")) if it.get("leaf") else None
        cost_only = it.get("model") != "build" and not it.get("uplift")
        saving = it.get("model") == "build" and it.get("kind") == "saving"
        if "ansoff" not in fw:
            if saving or cost_only:
                fw["ansoff"] = "efficiency"
            elif lf:
                fw["ansoff"] = "penetration"
            else:
                untagged.setdefault("ansoff", []).append(it["id"])
        if "four_p" not in fw:
            if saving or cost_only:
                fw["four_p"] = "none"
            elif lf and leaf_kind(lf, cur) == "price":
                fw["four_p"] = "price"
            else:
                untagged.setdefault("four_p", []).append(it["id"])
        if "value_chain" not in fw:
            untagged.setdefault("value_chain", []).append(it["id"])
        for lens, cats in LENS_TAGS.items():
            if fw.get(lens) is not None and fw[lens] not in cats:
                raise SpecError(f"initiative {it['id']}: frameworks.{lens} must be one of {', '.join(cats)}")
    for lens, ids in untagged.items():
        warn.append(f"frameworks.{lens} not set for {', '.join(ids)} - the {lens} view shows them as not tagged "
                    f"(set it in the specification card)")
    for lf in leaves.values():
        lf["_kind"] = leaf_kind(lf, cur)
    return used


class ObjectivesMixin:
    """Builder methods: Calculations section 10, the Objectives and Frameworks sheets, their Checks rows."""

    # ------------------------------------------------------------------ Calculations section 10
    def qty_row(self, it_id, q):
        return {"volume": lambda: self.buildrows[it_id]["vol"], "revenue": lambda: self.initrev.get(it_id),
                "saving": lambda: self.saverow.get(it_id), "opex": lambda: self.opexrow.get(it_id),
                "capex": lambda: self.capexrow.get(it_id), "ebit": lambda: self.initebit.get(it_id)}[q]()

    def build_objectives_calc(self):
        """Every KPI: base case (do nothing), effect of each strategy, with strategies, change."""
        self.kpis = []
        fy = self.years[-1]
        lc, fc = L(self.last_col), L(self.col(self.proj_years[0]))
        bc = L(self.base_col)
        self.csec("10. Case objectives: each KPI in the base case (do nothing) and with the strategies")
        inits = self.s["_initiatives"]
        rev = self.revenue_actual()
        ebitda = next((a for a in self.s.get("actuals", []) if a.get("role") == "ebitda"), None)
        for o in self.s["objectives"]:
            if o["kind"] == "profitability":
                ebit = self.rows["ebit"]
                if ebitda and rev:
                    m = self.put_scalar(f"Base-year EBITDA margin (reported, {self.base_year})", "%",
                                        f"=Actuals!{bc}{self.actrow[ebitda['id']]}/Actuals!{bc}{self.actrow[rev['id']]}",
                                        "obj_margin0")
                    b = self.put_row("EBITDA, base case (do nothing: base-year margin on baseline revenue)", self.cur,
                                     lambda c, p, y: f"={c}{self.rows['base_total']}*$E${m}", cols="base+")
                    w = self.put_row("EBITDA, with strategies (base case + incremental EBIT before new "
                                     "depreciation)", self.cur,
                                     lambda c, p, y: f"={c}{b}" if y == self.base_year else f"={c}{b}+{c}{ebit}",
                                     cols="base+")
                    name = "EBITDA"
                else:
                    b = self.put_row("Operating profit uplift, base case (do nothing)", self.cur,
                                     lambda c, p, y: "=0", cols="base+")
                    w = self.put_row("Operating profit uplift, with strategies (incremental EBIT)", self.cur,
                                     lambda c, p, y: f"={c}{b}" if y == self.base_year else f"={c}{ebit}",
                                     cols="base+")
                    name = "Operating profit uplift (incremental EBIT)"
                d = self.put_row(f"Change vs base case: {name}", self.cur, lambda c, p, y: f"={c}{w}-{c}{b}",
                                 cols="base+")
                eff = {it["id"]: [self.initebit[it["id"]]] for it in inits}
                self.kpis.append(dict(id="ebitda", name=name, unit=self.cur, obj=o, direction="higher",
                                      base=b, with_=w, change=d, effects=eff, ratio=False))
                if ebitda and rev:
                    mb = self.put_row("EBITDA margin, base case", "%",
                                      lambda c, p, y: f"={c}{b}/{c}{self.rows['base_total']}", cols="base+")
                    mw = self.put_row("EBITDA margin, with strategies", "%",
                                      lambda c, p, y: f"={c}{w}/{c}{self.rows['target_total']}", cols="base+")
                    md = self.put_row("Change vs base case: EBITDA margin (percentage points)", "%",
                                      lambda c, p, y: f"={c}{mw}-{c}{mb}", cols="base+")
                    self.kpis.append(dict(id="margin", name="EBITDA margin", unit="%", obj=o, direction="higher",
                                          base=mb, with_=mw, change=md, effects={}, ratio=True))
                self.blank()
            elif o["kind"] == "growth":
                d = self.put_row("Change vs base case: revenue", self.cur,
                                 lambda c, p, y: f"={c}{self.rows['target_total']}-{c}{self.rows['base_total']}",
                                 cols="base+")
                eff = {it["id"]: [self.initrev[it["id"]]] for it in inits if it["id"] in self.initrev}
                self.kpis.append(dict(id="revenue", name="Revenue", unit=self.cur, obj=o, direction="higher",
                                      base=self.rows["base_total"], with_=self.rows["target_total"], change=d,
                                      effects=eff, ratio=False))
                self.blank()
            else:
                for k in o["kpis"]:
                    ids = k.get("baseline") or []
                    g = self.live[k["baseline_growth"]] if k.get("baseline_growth") else None
                    lvl = "*".join(self.live[x] for x in ids)
                    gr = f"*(1+{g})^{{c}}$5" if g else ""
                    b = self.put_row(f"{k['name']}, base case (do nothing)", k["unit"],
                                     lambda c, p, y, lvl=lvl, gr=gr: f"={lvl}{gr.format(c=c)}" if lvl
                                     else "=0", cols="base+", ident=k["id"])
                    eff = {}
                    for e in k["effects"]:
                        it = next(x for x in inits if x["id"] == e["initiative"])
                        qr = self.qty_row(it["id"], e["quantity"])
                        mult = "".join(f"*{self.live[x]}" for x in e.get("factor") or []) + \
                            "".join(f"/{self.live[x]}" for x in e.get("per") or [])
                        sg = "-" if e.get("sign", 1) == -1 else ""
                        r = self.put_row(f"Effect on {k['name']}: {it['name']} ({QUANTITIES[e['quantity']]}"
                                         + (" x factor" if mult else "") + ")", k["unit"],
                                         lambda c, p, y, qr=qr, mult=mult, sg=sg: f"={sg}{c}{qr}{mult}",
                                         ident=it["id"])
                        eff.setdefault(it["id"], []).append(r)
                    rows_ = [r for rs in eff.values() for r in rs]
                    w = self.put_row(f"{k['name']}, with strategies", k["unit"],
                                     lambda c, p, y, b=b, rows_=rows_: f"={c}{b}" if y == self.base_year else
                                     f"={c}{b}+" + "+".join(f"{c}{r}" for r in rows_), cols="base+", ident=k["id"])
                    d = self.put_row(f"Change vs base case: {k['name']}", k["unit"],
                                     lambda c, p, y, b=b, w=w: f"={c}{w}-{c}{b}", cols="base+", ident=k["id"])
                    self.kpis.append(dict(id=k["id"], name=k["name"], unit=k["unit"], obj=o,
                                          direction=k["direction"], base=b, with_=w, change=d, effects=eff,
                                          ratio=False, stock=bool(k.get("stock")), spec=k))
                    self.blank()
        for k in self.kpis:
            k["ref_base"] = self.ref_of(self.put_scalar(f"KPI {fy}, base case: {k['name']}", k["unit"],
                                                        f"={lc}{k['base']}", f"kpi_b_{k['id']}"))
            k["ref_with"] = self.ref_of(self.put_scalar(f"KPI {fy}, with strategies: {k['name']}", k["unit"],
                                                        f"={lc}{k['with_']}", f"kpi_w_{k['id']}"))
            k["ref_chg"] = self.ref_of(self.put_scalar(f"KPI change in {fy} vs base case: {k['name']}", k["unit"],
                                                       f"={lc}{k['change']}", f"kpi_d_{k['id']}"))
            if not k["ratio"] and not k.get("stock"):
                k["ref_cum"] = self.ref_of(self.put_scalar(
                    f"KPI cumulative change {self.proj_years[0]}-{fy} vs base case: {k['name']}", k["unit"],
                    f"=SUM({fc}{k['change']}:{lc}{k['change']})", f"kpi_c_{k['id']}"))
        self.blank()

    def ref_of(self, r):
        return f"Calculations!$E${r}"

    # ------------------------------------------------------------------ helpers for the views
    def cell(self, row, col=None):
        return f"Calculations!{col or L(self.last_col)}{row}"

    def contrib(self, k, it_id):
        """Formula terms: the final-year change of KPI k caused by initiative it_id."""
        return [self.cell(r) for r in k["effects"].get(it_id, [])]

    def head(self, ws, r, labels, widths=None):
        for j, h in enumerate(labels, 2):
            c = ws.cell(r, j, h)
            c.font, c.fill = self.HEAD_FONT, self.HEAD_FILL
        return r + 1

    def put_terms(self, ws, r, c, terms, fmt):
        x = ws.cell(r, c, "=" + "+".join(terms) if terms else "=0")
        x.number_format = fmt
        return x

    # ------------------------------------------------------------------ Objectives sheet
    def build_objectives_sheet(self):
        ws = self.sh["Objectives"]
        fy = self.years[-1]
        self.title(ws, "Case objectives - how the strategies move them",
                   f"Base case = do nothing. Change = with strategies minus base case, in {fy} and summed over "
                   f"{self.proj_years[0]}-{fy}. All cells link to Calculations section 10.")
        for col_, w in zip("BCDEFGHIJ", (34, 30, 12, 15, 15, 15, 12, 18, 10)):
            ws.column_dimensions[col_].width = w
        r = self.head(ws, 4, ["Objective", "KPI", "Unit", f"Base case {fy}", f"With strategies {fy}",
                              f"Change {fy}", "Change %", f"Cumulative change {self.proj_years[0]}-{fy}",
                              "Better?"])
        self.objrow = {}
        for k in self.kpis:
            fmt = "0.0%" if k["unit"].startswith("%") else "#,##0.0"
            ws.cell(r, 2, f"{TBL[k['obj']['kind']]}: {k['obj']['name']}")
            ws.cell(r, 3, k["name"])
            ws.cell(r, 4, "pp" if k["ratio"] else k["unit"])
            for j, ref in ((5, k["ref_base"]), (6, k["ref_with"]), (7, k["ref_chg"])):
                ws.cell(r, j, f"={ref}").number_format = fmt
            ws.cell(r, 8, "" if k["ratio"] else f'=IF(E{r}=0,"n/a",G{r}/ABS(E{r}))').number_format = "0.0%"
            if k.get("stock"):
                ws.cell(r, 9, "a level, not summed")
            if not k["ratio"] and not k.get("stock"):
                ws.cell(r, 9, f"={k['ref_cum']}").number_format = fmt
            better = f"G{r}>0" if k["direction"] == "higher" else f"G{r}<0"
            ws.cell(r, 10, f'=IF(G{r}=0,"no change",IF({better},"better","worse"))')
            for j in (5, 6, 7):
                ws.cell(r, j).fill = self.OUT_FILL
            self.objrow[k["id"]] = r
            r += 1
        flows = [k for k in self.kpis if not k["ratio"]]
        r += 1
        ws.cell(r, 2, f"Contribution of each strategy to the change in {fy} (vs base case)").font = self.BOLD
        r = self.head(ws, r + 1, ["Pillar", "Strategy"] + [f"{k['name']} ({k['unit']})" for k in flows])
        first = r
        self.obj_init = {}
        for pl in self.s["pillars"]:
            for it in [x for x in self.s["_initiatives"] if x["pillar"] == pl["id"]]:
                ws.cell(r, 2, pl["name"])
                ws.cell(r, 3, it["name"])
                for j, k in enumerate(flows, 4):
                    self.put_terms(ws, r, j, self.contrib(k, it["id"]), "#,##0.0")
                self.obj_init[it["id"]] = r
                r += 1
        ws.cell(r, 2, "Total").font = self.BOLD
        for j, k in enumerate(flows, 4):
            x = ws.cell(r, j, f"=SUM({L(j)}{first}:{L(j)}{r - 1})")
            x.number_format, x.font = "#,##0.0", self.BOLD
        self.obj_total = (r, flows)
        r += 2
        ws.cell(r, 2, f"By pillar (change in {fy} vs base case)").font = self.BOLD
        r = self.head(ws, r + 1, ["Pillar", ""] + [f"{k['name']} ({k['unit']})" for k in flows])
        pfirst = r
        for pl in self.s["pillars"]:
            ws.cell(r, 2, pl["name"])
            rows_ = [self.obj_init[it["id"]] for it in self.s["_initiatives"] if it["pillar"] == pl["id"]]
            for j, _ in enumerate(flows, 4):
                ws.cell(r, j, "=" + "+".join(f"{L(j)}{x}" for x in rows_) if rows_ else "=0").number_format = \
                    "#,##0.0"
            r += 1
        ws.cell(r, 2, "Total").font = self.BOLD
        for j, _ in enumerate(flows, 4):
            ws.cell(r, j, f"=SUM({L(j)}{pfirst}:{L(j)}{r - 1})").number_format = "#,##0.0"
        self.obj_pillar_total = r
        if flows:
            bar = BarChart()
            bar.type, bar.title = "bar", f"{flows[0]['name']}: change in {fy} by strategy ({flows[0]['unit']})"
            bar.height, bar.width = 7.5, 16
            bar.add_data(Reference(ws, min_col=4, min_row=first, max_row=self.obj_total[0] - 1),
                         titles_from_data=False)
            bar.set_categories(Reference(ws, min_col=3, min_row=first, max_row=self.obj_total[0] - 1))
            bar.legend = None
            bar.x_axis.delete = bar.y_axis.delete = False
            ws.add_chart(bar, f"L4")

    # ------------------------------------------------------------------ Frameworks sheet
    def build_frameworks_sheet(self):
        ws = self.sh["Frameworks"]
        fy = self.years[-1]
        self.title(ws, "The impact through the consulting frameworks",
                   f"Each view splits the same change vs the base case ({fy}; revenue summed over "
                   f"{self.proj_years[0]}-{fy}); every view adds up to the strategy total (Checks).")
        ws.column_dimensions["B"].width = 52
        for j in range(3, 12):
            ws.column_dimensions[L(j)].width = 16
        ws.column_dimensions[L(12)].width = 60
        inits = self.s["_initiatives"]
        flows = [k for k in self.kpis if not k["ratio"]]
        prof = next((k for k in flows if k["id"] == "ebitda"), None)
        fc, lc = L(self.col(self.proj_years[0])), L(self.last_col)

        def cumrev(it_id):
            r = self.initrev.get(it_id)
            return [f"SUM(Calculations!{fc}{r}:{lc}{r})"] if r else []
        metrics = [("Cumulative incremental revenue", self.cur, cumrev)] + \
            [(f"{TBL[k['obj']['kind']]}: {k['name']} {fy}", k["unit"],
              (lambda it_id, k=k: self.contrib(k, it_id))) for k in flows]
        self.lens_totals = []
        r = 4

        def lens(title, note, cats, cat_of):
            nonlocal r
            ws.cell(r, 2, title).font = Font(bold=True, size=12)
            ws.cell(r + 1, 2, note).font = Font(italic=True, color="595959")
            r = self.head(ws, r + 2, ["Category"] + [f"{n} ({u})" for n, u, _ in metrics] + [""] * 0)
            ws.cell(r - 1, 3 + len(metrics), "Strategies").font = self.HEAD_FONT
            ws.cell(r - 1, 3 + len(metrics)).fill = self.HEAD_FILL
            first = r
            for key, label in cats:
                its = [it for it in inits if cat_of(it) == key]
                if not its and key == NOT_TAGGED:
                    continue
                ws.cell(r, 2, label)
                for j, (_, _, fn) in enumerate(metrics, 3):
                    self.put_terms(ws, r, j, [t for it in its for t in fn(it["id"])], "#,##0.0")
                ws.cell(r, 3 + len(metrics), ", ".join(it["name"] for it in its) or "-")
                r += 1
            ws.cell(r, 2, "Total").font = self.BOLD
            for j in range(3, 3 + len(metrics)):
                ws.cell(r, j, f"=SUM({L(j)}{first}:{L(j)}{r - 1})").number_format = "#,##0.0"
            self.lens_totals.append((title, r))
            r += 3

        # 1. Triple bottom line by pillar (the case objectives, pillar by pillar)
        lens("1. Triple bottom line (Elkington): profit, planet and people by pillar",
             "Rows are the team's pillars; columns are the case objectives.",
             [(pl["id"], pl["name"]) for pl in self.s["pillars"]], lambda it: it["pillar"])
        # 2. Profitability tree
        if prof:
            ws.cell(r, 2, f"2. Profitability tree: where the change in {prof['name']} comes from ({fy})").font = \
                Font(bold=True, size=12)
            ws.cell(r + 1, 2, "Profit = (volume x price) x margin - costs. Each strategy's effect, split by the "
                              "lever it pulls.").font = Font(italic=True, color="595959")
            r = self.head(ws, r + 2, ["Lever", f"{prof['name']} change {fy} ({self.cur})", "Strategies"])
            first = r
            m = self.live[self.s["finance"]["margin"]]
            parts = {"volume": [], "price": [], "new": [], "saving": [], "opex": []}
            who = {k: [] for k in parts}
            for it in inits:
                mi = self.live[it["margin"]] if it.get("margin") else m
                if it["id"] in self.initrev:
                    lf = self.s["_leaves"].get(it.get("leaf")) if it.get("leaf") else None
                    key = "new" if not lf else ("price" if lf["_kind"] == "price" else "volume")
                    parts[key].append(f"{self.cell(self.initrev[it['id']])}*Calculations!{mi}")
                    who[key].append(it["name"])
                if it["id"] in self.saverow:
                    parts["saving"].append(self.cell(self.saverow[it["id"]]))
                    who["saving"].append(it["name"])
                if it["id"] in self.opexrow:
                    parts["opex"].append(f"-{self.cell(self.opexrow[it['id']])}")
                    who["opex"].append(it["name"])
            for key, label in (("volume", "Revenue: more volume (customers, visits, occupancy), after variable costs"),
                               ("price", "Revenue: higher price or spend per unit, after variable costs"),
                               ("new", "Revenue: new streams (volume x price of a new offer), after variable costs"),
                               ("saving", "Cost: savings (cost avoided)"),
                               ("opex", "Cost: operating cost added by the strategies")):
                ws.cell(r, 2, label)
                x = ws.cell(r, 3, ("=" + "".join(t if t.startswith("-") else "+" + t for t in parts[key]).lstrip("+"))
                            if parts[key] else "=0")
                x.number_format = "#,##0.0"
                ws.cell(r, 4, ", ".join(who[key]) or "-")
                r += 1
            ws.cell(r, 2, "Total").font = self.BOLD
            ws.cell(r, 3, f"=SUM(C{first}:C{r - 1})").number_format = "#,##0.0"
            self.prof_tree_total = r
            r += 3
        # 3-5. tagged lenses
        for n, (lens_id, title, note) in enumerate((
                ("ansoff", "Ansoff matrix: growth vector", "Where the growth comes from: existing or new products "
                                                           "and markets; efficiency moves cost, not growth."),
                ("four_p", "Marketing mix (4Ps)", "Which marketing lever each strategy pulls."),
                ("value_chain", "Value chain (Porter)", "Which activity of the business each strategy changes.")),
                3 if prof else 2):
            cats = list(LENS_TAGS[lens_id].items()) + [(NOT_TAGGED, NOT_TAGGED)]
            lens(f"{n}. {title}", note, cats,
                 lambda it, lens_id=lens_id: it["frameworks"].get(lens_id) or NOT_TAGGED)
        # balanced scorecard (Kaplan & Norton, with the triple-bottom-line outcomes)
        n += 1
        ws.cell(r, 2, f"{n}. Balanced scorecard: outcomes, customers, processes, enablers").font = \
            Font(bold=True, size=12)
        ws.cell(r + 1, 2, "Kaplan & Norton's four perspectives; the financial perspective widened to the case "
                          "objectives.").font = Font(italic=True, color="595959")
        r = self.head(ws, r + 2, ["Perspective", "Measure", "Value", "Unit"])
        cum = lambda row: f"SUM(Calculations!{fc}{row}:{lc}{row})"
        rows_ = []
        for k in self.kpis:
            rows_.append(("Outcomes (case objectives)", f"{k['name']}: change in {fy} vs base case",
                          f"={k['ref_chg']}", "pp" if k["ratio"] else k["unit"]))
        rows_.append(("Outcomes (case objectives)", "Valuation: NPV of the strategies", f"={self.ref['npv']}",
                      self.cur))
        vol = [self.initrev[it["id"]] for it in inits if it["id"] in self.initrev and it.get("leaf")
               and self.s["_leaves"][it["leaf"]]["_kind"] == "volume"]
        new = [self.initrev[it["id"]] for it in inits if it["id"] in self.initrev and not it.get("leaf")]
        rows_.append(("Customers", f"Revenue from more customers or use, {self.proj_years[0]}-{fy}",
                      "=" + "+".join(cum(x) for x in vol) if vol else "=0", self.cur))
        rows_.append(("Customers", f"Revenue from new offers, {self.proj_years[0]}-{fy}",
                      "=" + "+".join(cum(x) for x in new) if new else "=0", self.cur))
        rows_.append(("Internal processes", f"Cost savings, {self.proj_years[0]}-{fy}",
                      "=" + "+".join(cum(x) for x in self.saverow.values()) if self.saverow else "=0", self.cur))
        rows_.append(("Learning & growth (enablers)", f"Capex invested, {self.proj_years[0]}-{fy}",
                      f"={cum(self.rows['capex'])}" if "capex" in self.rows else "=0", self.cur))
        rows_.append(("Learning & growth (enablers)", f"Operating cost of running the strategies, "
                                                      f"{self.proj_years[0]}-{fy}",
                      f"={cum(self.rows['opex'])}" if "opex" in self.rows else "=0", self.cur))
        for a, b, f, u in rows_:
            ws.cell(r, 2, a)
            ws.cell(r, 3, b)
            ws.cell(r, 4, f).number_format = "0.0%" if u in ("%", "pp") else "#,##0.0"
            ws.cell(r, 5, u)
            r += 1

    def objectives_checks(self, row):
        """Rows for the model's own Checks sheet: every view adds up to the same total."""
        tol = f"Inputs!$F${self.inrow['sum_tol']}"
        tr, flows = self.obj_total
        for j, k in enumerate(flows, 4):
            row(f"Objectives: strategies add up to the change in {k['name']}",
                f"=ABS(Objectives!{L(j)}{tr}-{k['ref_chg']})", f"={tol}", "C{r}<=D{r}")
            row(f"Objectives: pillars add up to the change in {k['name']}",
                f"=ABS(Objectives!{L(j)}{self.obj_pillar_total}-{k['ref_chg']})", f"={tol}", "C{r}<=D{r}")
        prof = next((k for k in flows if k["id"] == "ebitda"), None)
        for title, r_ in self.lens_totals:
            for j, k in enumerate(flows, 4):
                row(f"Frameworks: '{title.split(':')[0]}' adds up to the change in {k['name']}",
                    f"=ABS(Frameworks!{L(j)}{r_}-{k['ref_chg']})", f"={tol}", "C{r}<=D{r}")
        if prof and hasattr(self, "prof_tree_total"):
            row(f"Frameworks: profitability tree adds up to the change in {prof['name']}",
                f"=ABS(Frameworks!C{self.prof_tree_total}-{prof['ref_chg']})", f"={tol}", "C{r}<=D{r}")
