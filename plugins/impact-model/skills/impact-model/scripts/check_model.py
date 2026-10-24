#!/usr/bin/env python3
"""Independent integrity checker for financial impact models (.xlsx).

The model must not grade its own work: this script opens the workbook,
recalculates it, re-runs every Data Table by substitution, and reports
PASS / FAIL per check.

    python scripts/check_model.py model.xlsx
    python scripts/check_model.py model.xlsx --json report.json
    python scripts/check_model.py model.xlsx --trace "Outputs!E8"   # source chain of one cell

Recalculation engine (--engine auto): LibreOffice headless if installed,
otherwise the `formulas` library, otherwise the built-in engine (xleval.py).
The built-in engine is always used for what-if work (Data Tables, axis tests)
and is cross-checked against the main engine.

Exit code 0 only when every check passes.
Requires: openpyxl. Optional: LibreOffice (soffice) or `pip install formulas`.
"""
from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from openpyxl.formula.tokenizer import Token, Tokenizer  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from xleval import XLErr, Model, addr, is_num, parse_cell, run_big_stack, split_sheet  # noqa: E402

# Default cap on year-on-year growth of level rows (revenue, customers, price...).
# 25%/yr is roughly the top-decile organic growth of established firms; anything
# faster needs a written flag ("FLAG: ..." in the row) explaining why.
GROWTH_CAP = 0.25
REL_TOL = 1e-6          # totals must match their parts to 1 part in a million
MAX_LIST = 12           # findings printed per check
HIST_TOL = 0.01         # baseline growth may sit 1 pp outside the historical band before it needs a justification
HIST_Q = (0.25, 0.75)   # the band: interquartile range of reported YoY revenue growth, widened to include the CAGR

ROLE_SKIP = re.compile(r"cover|readme|read me|chart|list|note|request|tree|instruc|about|log", re.I)
ROLE_INPUT = re.compile(r"input|assum|actual|financials|historic|data|source|register|p&l|peer", re.I)
ROLE_OUTPUT = re.compile(r"output|objective|framework|sensitiv|dashboard|bridge|summary|scenario|tornado|result", re.I)
ROLE_CHECK = re.compile(r"^checks?$", re.I)

LBL_TOTAL = re.compile(r"\b(sub)?total\b", re.I)
LBL_RATIO = re.compile(r"%|growth|margin|\brate\b|\bratios?\b|factor|multiplier|discount|share|penetration|"
                       r"contribution|per cent|percent|wacc|irr|elasticity|adoption|probabil|\blog\b|change factor", re.I)
LBL_PEN = re.compile(r"penetration|market share|share of|% of (the )?(market|addressable|capacity|ceiling)|"
                     r"occupancy|utili[sz]ation|attach rate|adoption|load factor|% of ceiling", re.I)
LBL_VOLUME = re.compile(r"\b(fans|attendance|visitors|customers|users|subscribers|members|visits|volumes?|"
                        r"units|guests|passengers|students|patients|clients|households|accounts|footfall)\b", re.I)
LBL_SKIP_GROWTH = re.compile(r"incremental|uplift|impact|increase|delta|additional|\bcost|capex|opex|cash|"
                             r"\bpv\b|npv|cumulative|payback|flag|check|period|index|discount|year|tax|"
                             r"strateg|initiative|saving", re.I)
LBL_DECAY = re.compile(r"flatten|decay|fade|erosion|churn|attrition", re.I)

PERIOD_RE = re.compile(r"^\s*(?:FY\s*)?(\d{4})\s*([AEFBP])?\s*(?:[-\s]?(H|Q)\s*([1-4]))?\s*$", re.I)
PERIOD_RE2 = re.compile(r"^\s*(H|Q)([1-4])\s*(?:FY)?\s*(\d{4})\s*([AEFBP])?\s*$", re.I)
FISCAL_RE = re.compile(r"^\s*(\d{4})\s*/\s*(\d{2,4})\s*$")


# ----------------------------------------------------------------------------- recalculation
def find_soffice():
    for name in ("soffice", "libreoffice"):
        p = shutil.which(name)
        if p:
            return p
    for p in ("/Applications/LibreOffice.app/Contents/MacOS/soffice",
              "/usr/lib/libreoffice/program/soffice", "/opt/libreoffice/program/soffice"):
        if os.path.exists(p):
            return p
    return None


XCU = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema">
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="ODFRecalcMode" oor:op="fuse"><value>0</value></prop></item>
</oor:items>
"""


def recalc_libreoffice(path, soffice):
    """Return path of a LibreOffice-recalculated copy (never touches the original)."""
    tmp = tempfile.mkdtemp(prefix="imcheck_lo_")
    prof = os.path.join(tmp, "profile")
    os.makedirs(os.path.join(prof, "user"), exist_ok=True)
    with open(os.path.join(prof, "user", "registrymodifications.xcu"), "w") as f:
        f.write(XCU)  # force "always recalculate on load" for xlsx
    src = os.path.join(tmp, "in_" + os.path.basename(path))
    shutil.copy(path, src)
    out = os.path.join(tmp, "out")
    os.makedirs(out)
    cmd = [soffice, f"-env:UserInstallation=file://{prof}", "--headless", "--norestore",
           "--convert-to", "xlsx", "--outdir", out, src]
    subprocess.run(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=240, check=False)
    res = os.path.join(out, os.path.basename(src))
    if not os.path.exists(res):
        raise RuntimeError("LibreOffice did not produce a recalculated file")
    return res


def recalc_formulas(path, model):
    """Recalculate with the `formulas` library on a copy without Data Tables."""
    import contextlib
    import io
    import openpyxl
    with contextlib.redirect_stderr(io.StringIO()):
        import formulas  # noqa
    from openpyxl.worksheet.formula import DataTableFormula
    tmp = tempfile.mkdtemp(prefix="imcheck_fx_")
    wb = openpyxl.load_workbook(path)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, DataTableFormula):
                    c.value = None
    cp = os.path.join(tmp, "copy.xlsx")
    wb.save(cp)
    with contextlib.redirect_stderr(io.StringIO()), contextlib.redirect_stdout(io.StringIO()):
        xl = formulas.ExcelModel().loads(cp).finish()
        sol = xl.calculate()
    upper = {s.upper(): s for s in model.sheetnames}
    vals = {}
    for k, v in sol.items():
        m = re.match(r"^'?\[[^\]]+\](.+?)'?!([A-Z]+\d+)$", str(k))
        if not m:
            continue
        sheet = upper.get(m.group(1).upper())
        rc = parse_cell(m.group(2))
        if not sheet or not rc:
            continue
        try:
            x = v.value[0][0]
        except Exception:
            x = v
        if hasattr(x, "item"):
            x = x.item()
        if type(x).__name__ == "XlError":
            x = XLErr(str(x))
        if isinstance(x, str) and x.startswith("#"):
            x = XLErr(x)
        vals[(sheet, rc[0], rc[1])] = x
    return vals


class Values:
    """Recalculated values for every cell + the built-in what-if context."""

    def __init__(self, path, engine="auto"):
        self.model = Model(path)
        self.ctx = self.model.context()
        self.engine = "built-in (xleval)"
        self.notes = []
        self.engine_vals = None
        soffice = find_soffice() if engine in ("auto", "libreoffice") else None
        if soffice:
            try:
                rc = recalc_libreoffice(path, soffice)
                lo = Model(path, cached_path=rc)
                self.engine_vals = {k: v for k, v in lo.cached.items() if k in self.model.cells}
                self.engine = "LibreOffice headless"
            except Exception as e:  # noqa
                self.notes.append(f"LibreOffice recalculation failed ({e}); falling back")
        elif engine == "libreoffice":
            self.notes.append("LibreOffice not found; falling back")
        if self.engine_vals is None and engine in ("auto", "formulas", "libreoffice"):
            try:
                self.engine_vals = run_big_stack(recalc_formulas, path, self.model)
                self.engine = "formulas library"
            except ImportError:
                self.notes.append("`formulas` not installed (pip install formulas) - using built-in engine only")
            except Exception as e:  # noqa
                self.notes.append(f"`formulas` could not evaluate this workbook ({type(e).__name__}); using built-in engine")
        self.builtin = run_big_stack(self._eval_all)
        self.dt_dependent = set()
        # Neither `formulas` (strips them) nor LibreOffice (skips tables saved without cached values)
        # reliably recomputes Data Tables: they and everything downstream come from the built-in engine.
        if self.engine_vals is not None and self.model.datatables:
            dep = self.model.dependents()
            stack = list(self.model.dt_index)
            while stack:
                k = stack.pop()
                for d in dep.get(k, ()):
                    if d not in self.dt_dependent:
                        self.dt_dependent.add(d)
                        stack.append(d)
        self.diffs = []
        if self.engine_vals:
            for k, v in self.builtin.items():
                if k in self.model.dt_index or k in self.dt_dependent:
                    continue
                e = self.engine_vals.get(k)
                if e is None:
                    continue
                if is_num(v) and is_num(e):
                    if abs(v - e) > 1e-6 * max(1.0, abs(e)):
                        self.diffs.append((k, v, e))
                elif is_num(v) != is_num(e) and not (isinstance(v, XLErr) and isinstance(e, XLErr)):
                    self.diffs.append((k, v, e))
        self.stale = []
        for k, v in self.builtin.items():
            c = self.model.cached.get(k)
            if is_num(c) and is_num(v) and abs(v - c) > 1e-6 * max(1.0, abs(c)):
                self.stale.append((k, v, c))

    def _eval_all(self):
        out = {}
        for k in list(self.model.cells):
            if self.model.is_formula(k) or k in self.model.dt_index:
                out[k] = self.ctx.get(k)
        for k in self.model.dt_index:
            out.setdefault(k, self.ctx.get(k))
        return out

    def get(self, key):
        """Best recalculated value for a cell."""
        m = self.model
        if key in m.dt_index or key in self.dt_dependent:
            return self.builtin.get(key)
        if m.is_formula(key):
            if self.engine_vals is not None and key in self.engine_vals and self.engine != "built-in (xleval)":
                v = self.engine_vals[key]
                return v
            return self.builtin.get(key)
        v = m.cells.get(key)
        return None if not (is_num(v) or isinstance(v, str) or isinstance(v, bool)) else v


# ----------------------------------------------------------------------------- workbook reading helpers
class Book:
    def __init__(self, vals: Values):
        self.v = vals
        self.m = vals.model
        self.roles = {s: self.role(s) for s in self.m.sheetnames}
        self._labels = {}
        self.timelines = self.find_timelines()
        self.registers = self.find_registers()

    @staticmethod
    def role(name):
        if ROLE_CHECK.search(name):
            return "check"
        if ROLE_SKIP.search(name):
            return "skip"
        if ROLE_OUTPUT.search(name):
            return "output"
        if ROLE_INPUT.search(name):
            return "input"
        return "calc"

    def sheets(self, *roles):
        return [s for s in self.m.sheetnames if self.roles[s] in roles]

    def val(self, key):
        return self.v.get(key)

    def text(self, key):
        x = self.val(key)
        return x.strip() if isinstance(x, str) else ""

    def label(self, key):
        """Nearest text to the left of a cell on the same row (the row's label)."""
        sh, r, c = key
        if (sh, r, c) in self._labels:
            return self._labels[(sh, r, c)]
        lab = ""
        for cc in range(c - 1, 0, -1):
            t = self.text((sh, r, cc))
            if t and not t.startswith("#") and not re.fullmatch(r"[A-F]\s*[1-6]", t):
                lab = t
                break
        self._labels[(sh, r, c)] = lab
        return lab

    def row_label(self, sh, r, upto=8):
        for c in range(1, upto + 1):
            t = self.text((sh, r, c))
            if t and not t.startswith("#"):
                return t
        return ""

    def number_format(self, key):
        try:
            return self.m.wb[key[0]].cell(key[1], key[2]).number_format or ""
        except (KeyError, ValueError):
            return ""

    def name(self, key):
        lab = self.label(key)
        return f"{key[0]}!{addr(key[1], key[2])}" + (f" ({lab})" if lab else "")

    def numeric_constant(self, key):
        if key in self.m.dt_index:
            return False  # Data Table results are computed by Excel, not typed
        x = self.m.cells.get(key)
        return is_num(x)

    # -- timelines
    @staticmethod
    def parse_period(x):
        if is_num(x) and float(x).is_integer() and 1990 <= x <= 2100:
            return float(x), None
        if not isinstance(x, str):
            return None
        m = PERIOD_RE.match(x)
        if m:
            y = float(m.group(1))
            if not 1990 <= y <= 2100:
                return None
            sub = 0.0
            if m.group(3):
                n = int(m.group(4))
                sub = (n - 1) * (0.5 if m.group(3).upper() == "H" else 0.25)
            return y + sub, (m.group(2) or "").upper() or None
        m = PERIOD_RE2.match(x)
        if m:
            n = int(m.group(2))
            sub = (n - 1) * (0.5 if m.group(1).upper() == "H" else 0.25)
            return float(m.group(3)) + sub, (m.group(4) or "").upper() or None
        m = FISCAL_RE.match(x)
        if m:
            return float(m.group(1)), None
        return None

    def find_timelines(self):
        """Per sheet: list of segments {row, cols:[...], periods:[...], markers:[...]}."""
        out = {}
        for sh in self.m.sheetnames:
            mr, mc = self.m.max_rc.get(sh, (0, 0))
            segs = []
            rows = {}
            for (s, r, c) in self.m.cells:
                if s == sh:
                    rows.setdefault(r, []).append(c)
            for r, cols in rows.items():
                cols = sorted(cols)
                cur = []
                for c in cols:
                    p = self.parse_period(self.val((sh, r, c)))
                    if p and (not cur or c == cur[-1][0] + 1):
                        cur.append((c, p))
                    else:
                        if len(cur) >= 3:
                            segs.append((r, cur))
                        cur = [(c, p)] if p else []
                if len(cur) >= 3:
                    segs.append((r, cur))
            out[sh] = [{"row": r, "cols": [c for c, _ in seg], "periods": [p[0] for _, p in seg],
                        "markers": [p[1] for _, p in seg]} for r, seg in segs]
            for t in out[sh]:
                if all(mk is None for mk in t["markers"]):
                    for u in out[sh]:
                        if u is not t and u["cols"] == t["cols"] and u["periods"] == t["periods"] \
                                and any(u["markers"]):
                            t["markers"] = list(u["markers"])
                            break
            # a data row of years is not a timeline unless it is mostly a header
            out[sh] = [t for t in out[sh] if not self._looks_like_data(sh, t)]
        return out

    def _looks_like_data(self, sh, t):
        steps = {round(b - a, 4) for a, b in zip(t["periods"], t["periods"][1:])}
        return any(s <= 0 for s in steps) and len(set(t["periods"])) < len(t["periods"]) / 2

    def timeline_for(self, sh, r):
        """Timeline header nearest above row r on sheet sh."""
        best = None
        for t in self.timelines.get(sh, []):
            if t["row"] < r and (best is None or t["row"] > best["row"]):
                best = t
        return best

    # -- registers (assumption tables with Source / Grade columns)
    def find_registers(self):
        regs = []
        for sh in self.m.sheetnames:
            if self.roles[sh] == "skip":
                continue
            for r in range(1, 16):
                heads = {}
                for c in range(1, 40):
                    t = self.text((sh, r, c)).lower()
                    if t:
                        heads[c] = t
                src = [c for c, t in heads.items() if t.startswith("source") or t == "source"]
                grd = [c for c, t in heads.items() if "grade" in t or "reliab" in t]
                if src and grd:
                    col = lambda pat: next((c for c, t in heads.items() if re.search(pat, t)), None)
                    regs.append({
                        "sheet": sh, "header": r, "source": src[0], "grade": grd[0],
                        "date": col(r"^date|as of|retrieved|vintage"),
                        "low": col(r"^low"), "base": col(r"^base|^value"), "high": col(r"^high"),
                        "type": col(r"^type|^kind"), "bench_low": col(r"bench.*low|ref.*low"),
                        "bench_high": col(r"bench.*high|ref.*high"), "justif": col(r"justif"),
                        "bench_med": col(r"bench.*med|ref.*med"), "bench_src": col(r"bench.*source|ref.*source"),
                        "link": col(r"link|path|url"), "id": col(r"^id$|^#$"),
                    })
                    break
        return regs

    def register_row(self, key):
        for reg in self.registers:
            if key[0] == reg["sheet"] and key[1] > reg["header"]:
                return reg
        return None

    def sourced(self, key):
        reg = self.register_row(key)
        if not reg:
            return False
        r = key[1]
        s = self.text((key[0], r, reg["source"]))
        g = self.text((key[0], r, reg["grade"]))
        return bool(s) and bool(re.fullmatch(r"[A-F]\s*[1-6]", g.upper()))

    def controls(self):
        """Cells that are legitimately typed on calc/output sheets: Data Table input cells and axes,
        and cells named ctl_* (scenario switch, tornado selector)."""
        ctl = set()
        for dt in self.m.datatables:
            for inp in (dt.in1, dt.in2):
                if inp:
                    ctl.add((dt.sheet,) + inp)
            top, left = dt.axis_cells()
            ctl |= {(dt.sheet,) + a for a in top + left}
        for name, txt in self.m.names.items():
            if name.lower().startswith("ctl_"):
                r = self.m.resolve(txt, None)
                if r:
                    ctl.add((r[0], r[1], r[2]))
        return ctl


# ----------------------------------------------------------------------------- checks
class Result:
    def __init__(self, cid, title):
        self.id, self.title = cid, title
        self.fails, self.warns, self.info = [], [], []

    @property
    def status(self):
        return "FAIL" if self.fails else ("WARN" if self.warns else "PASS")

    def as_dict(self):
        return {"id": self.id, "title": self.title, "status": self.status,
                "fails": self.fails, "warnings": self.warns, "info": self.info}


def same(a, b):
    return all(is_num(x) and is_num(y) and abs(x - y) <= 1e-9 * max(1.0, abs(x), abs(y)) for x, y in zip(a, b))


def dup_pairs(vectors, labels):
    out = []
    for i in range(len(vectors)):
        for j in range(i + 1, len(vectors)):
            if vectors[i] and len(vectors[i]) == len(vectors[j]) and same(vectors[i], vectors[j]):
                out.append((labels[i], labels[j]))
    return out


def fmt(x):
    if is_num(x):
        return f"{x:,.6g}"
    return repr(x)


def check_recalc(bk: Book):
    R = Result("recalc", "Recalculated; no formula errors")
    v = bk.v
    R.info.append(f"engine: {v.engine}; built-in engine used for Data Tables and what-if tests")
    R.info += v.notes
    if v.diffs:
        R.warns.append(f"{len(v.diffs)} cells differ between {v.engine} and the built-in engine, e.g. " +
                       "; ".join(f"{bk.name(k)} {fmt(a)} vs {fmt(b)}" for k, a, b in v.diffs[:5]))
    if v.stale:
        R.info.append(f"{len(v.stale)} cells differ from the values saved in the file (stale cache or Data Table "
                      f"artefacts), e.g. " + "; ".join(f"{bk.name(k)} saved {fmt(c)} vs recalculated {fmt(a)}"
                                                       for k, a, c in v.stale[:4]))
    errs = []
    for k in bk.m.cells:
        if bk.roles.get(k[0]) == "skip":
            continue
        if bk.m.is_formula(k) or k in bk.m.dt_index:
            x = v.get(k)
            if isinstance(x, XLErr) or (isinstance(x, str) and x.startswith("#") and x.endswith(("!", "?", "A"))):
                errs.append(f"{bk.name(k)} = {x}")
    if errs:
        R.fails.append(f"{len(errs)} formula errors: " + "; ".join(errs[:MAX_LIST]))
    unsup = bk.v.ctx.unsupported
    if unsup:
        R.warns.append("built-in engine does not support: " + ", ".join(sorted(unsup)))
    return R


def find_pasted_grids(bk: Book):
    """Blocks of typed numbers laid out like a sensitivity table (formula corner + two axes)."""
    m = bk.m
    grids = []
    taken = set()
    for key in sorted(k for k in m.cells if m.is_formula(k) and bk.roles[k[0]] in ("calc", "output")):
        sh, r, c = key
        if (sh, r + 1, c + 1) in taken:
            continue
        num = lambda k: is_num(bk.val(k))
        if not all(num((sh, r, c + j)) for j in (1, 2, 3)) or not all(num((sh, r + i, c)) for i in (1, 2, 3)):
            continue
        if not all(bk.numeric_constant((sh, r + i, c + j)) and (sh, r + i, c + j) not in m.dt_index
                   for i in (1, 2, 3) for j in (1, 2, 3)):
            continue
        ncols = 0
        while bk.numeric_constant((sh, r + 1, c + 1 + ncols)):
            ncols += 1
        nrows = 0
        while bk.numeric_constant((sh, r + 1 + nrows, c + 1)):
            nrows += 1
        g = {"sheet": sh, "corner": key, "r1": r + 1, "c1": c + 1, "r2": r + nrows, "c2": c + ncols}
        for i in range(nrows):
            for j in range(ncols):
                taken.add((sh, r + 1 + i, c + 1 + j))
        grids.append(g)
    return grids


def check_sensitivity(bk: Book):
    R = Result("sensitivity", "Sensitivity tables: axes feed the output; rows and columns differ")
    m, ctx = bk.m, bk.v.ctx
    if not m.datatables and not find_pasted_grids(bk):
        R.warns.append("no sensitivity tables found")
    for dt in m.datatables:
        sh = dt.sheet
        tag = f"{sh}!{dt.ref}"
        grid = [[bk.v.builtin.get((sh, r, c)) for c in range(dt.c1, dt.c2 + 1)] for r in range(dt.r1, dt.r2 + 1)]
        cached = [[m.cached.get((sh, r, c)) for c in range(dt.c1, dt.c2 + 1)] for r in range(dt.r1, dt.r2 + 1)]
        has_cache = all(is_num(x) for row in cached for x in row)
        top, left = dt.axis_cells()
        topv = [ctx.get((sh,) + a) for a in top]
        leftv = [ctx.get((sh,) + a) for a in left]
        # 1. every axis input cell must be a precedent of the output formula(s)
        pairs = [(dt.in1, "row-axis" if dt.two_d or dt.row_oriented else "column-axis")]
        if dt.two_d:
            pairs.append((dt.in2, "column-axis"))
        formula_cells = {dt.plan(r, c)[0] for r in range(dt.r1, dt.r2 + 1) for c in range(dt.c1, dt.c2 + 1)}
        for inp, kind in pairs:
            if inp is None:
                continue
            ik = (sh,) + inp
            feeds = any(ik in m.closure((sh,) + f) for f in formula_cells)
            if not feeds:
                R.fails.append(f"{tag}: {kind} input {bk.name(ik)} does not feed the output "
                               f"{bk.name((sh,) + dt.corner()) if dt.two_d else 'formula'} - this axis changes nothing")
        # 2. axis headers must not depend on the input cell they drive
        for inp, axis in ((dt.in1, top if (dt.two_d or dt.row_oriented) else left), (dt.in2, left)):
            if inp is None or not dt.two_d and axis is left and inp == dt.in2:
                continue
            ik = (sh,) + inp
            for a in axis:
                ak = (sh,) + a
                if m.is_formula(ak) and (ik in m.closure(ak) or ak == ik):
                    R.fails.append(f"{tag}: axis cell {bk.name(ak)} = {m.formula(ak)} depends on the table's own "
                                   f"input cell {addr(*inp)}; Excel substitutes into it and repeats a row/column")
        # 3. axes ordered
        for nm, vals in (("top axis", topv), ("left axis", leftv)):
            nums = [x for x in vals if is_num(x)]
            if len(nums) >= 3:
                inc = all(b > a for a, b in zip(nums, nums[1:]))
                dec = all(b < a for a, b in zip(nums, nums[1:]))
                if not (inc or dec):
                    R.fails.append(f"{tag}: {nm} is not strictly ordered: {', '.join(fmt(x) for x in nums)}")
        # 4. rows and columns must differ
        rows_lbl = [f"row {fmt(x)}" for x in leftv] if left else [f"row {i + 1}" for i in range(len(grid))]
        cols_lbl = [f"col {fmt(x)}" for x in topv] if top else [f"col {j + 1}" for j in range(len(grid[0]))]
        cols = [[grid[i][j] for i in range(len(grid))] for j in range(len(grid[0]))]
        if dt.two_d:
            for what, d in (("rows", dup_pairs(grid, rows_lbl)), ("columns", dup_pairs(cols, cols_lbl))):
                if d:
                    R.fails.append(f"{tag}: recalculated {what} identical: " + "; ".join(f"{a} = {b}" for a, b in d[:4]))
            if has_cache:
                ccols = [[cached[i][j] for i in range(len(cached))] for j in range(len(cached[0]))]
                for what, d in (("rows", dup_pairs(cached, rows_lbl)), ("columns", dup_pairs(ccols, cols_lbl))):
                    if d:
                        R.fails.append(f"{tag}: rows/columns shown in the file are identical ({what}): " +
                                       "; ".join(f"{a} = {b}" for a, b in d[:4]))
        else:
            vecs = cols if not dt.row_oriented else grid
            for j, vec in enumerate(vecs):
                nums = [x for x in vec if is_num(x)]
                if len(nums) >= 2 and max(nums) - min(nums) <= 1e-9 * max(1.0, abs(max(nums))):
                    R.fails.append(f"{tag}: output {j + 1} does not change across the axis")
        bad = [x for row in grid for x in row if not is_num(x)]
        if bad:
            R.fails.append(f"{tag}: {len(bad)} cells do not evaluate to numbers (e.g. {bad[0]})")
        R.info.append(f"{tag}: Data Table {'2-way' if dt.two_d else '1-way'} recalculated by substitution")
    for g in find_pasted_grids(bk):
        sh = g["sheet"]
        tag = f"{sh}!{addr(g['r1'], g['c1'])}:{addr(g['r2'], g['c2'])}"
        R.fails.append(f"{tag}: sensitivity grid is typed/pasted numbers, not a live Data Table")
        grid = [[bk.val((sh, r, c)) for c in range(g["c1"], g["c2"] + 1)] for r in range(g["r1"], g["r2"] + 1)]
        leftv = [bk.val((sh, r, g["c1"] - 1)) for r in range(g["r1"], g["r2"] + 1)]
        topv = [bk.val((sh, g["r1"] - 1, c)) for c in range(g["c1"], g["c2"] + 1)]
        cols = [[grid[i][j] for i in range(len(grid))] for j in range(len(grid[0]))]
        for what, d in (("rows", dup_pairs(grid, [f"row {fmt(x)}" for x in leftv])),
                        ("columns", dup_pairs(cols, [f"col {fmt(x)}" for x in topv]))):
            if d:
                R.fails.append(f"{tag}: identical {what}: " + "; ".join(f"{a} = {b}" for a, b in d[:4]))
        for nm, vals in (("top axis", topv), ("left axis", leftv)):
            nums = [x for x in vals if is_num(x)]
            if len(nums) >= 3 and not (all(b > a for a, b in zip(nums, nums[1:])) or
                                       all(b < a for a, b in zip(nums, nums[1:]))):
                R.fails.append(f"{tag}: {nm} is not strictly ordered: {', '.join(fmt(x) for x in nums)}")
    return R


def check_totals(bk: Book):
    R = Result("totals", "Every total equals the sum of its parts")
    m = bk.m
    n = 0
    for sh in bk.sheets("calc", "output", "input", "check"):
        model_sheet = bk.roles[sh] != "input"
        tl_rows = {t["row"] for t in bk.timelines.get(sh, [])}
        rows = sorted({k[1] for k in m.cells if k[0] == sh})
        for r in rows:
            lab = bk.row_label(sh, r)
            if not LBL_TOTAL.search(lab) or LBL_RATIO.search(lab):
                continue
            cols = sorted(k[2] for k in m.cells if k[0] == sh and k[1] == r and is_num(bk.val(k)))
            for c in cols:
                if not model_sheet and not m.is_formula((sh, r, c)):
                    continue  # typed totals in raw data (e.g. balance-sheet "Total capital") are not sums
                tot = bk.val((sh, r, c))
                parts, ratio_rows = [], []
                rr = r - 1
                while rr >= 1 and is_num(bk.val((sh, rr, c))) and rr not in tl_rows:
                    plab = bk.row_label(sh, rr)
                    if LBL_TOTAL.search(plab):
                        break
                    x = bk.val((sh, rr, c))
                    if (LBL_RATIO.search(plab) and x != 0 and (abs(x) <= 1.5 or "%" in bk.number_format((sh, rr, c)))) or \
                            (abs(x) < 1 and abs(tot) > 100 and not plab):
                        ratio_rows.append(rr)
                    else:
                        parts.append(x)
                    rr -= 1
                if not parts:
                    continue
                n += 1
                s = sum(parts)
                if abs(tot - s) > max(REL_TOL * max(abs(tot), abs(s)), 1e-6):
                    R.fails.append(f"{bk.name((sh, r, c))} = {fmt(tot)} but its parts (rows {rr + 1}-{r - 1}) sum to "
                                   f"{fmt(s)} [{m.formula((sh, r, c)) or 'typed value'}]")
                key = (sh, r, c)
                if ratio_rows and m.is_formula(key):
                    refs = m.precedents().get(key, set())
                    hit = [x for x in ratio_rows if (sh, x, c) in refs]
                    if hit:
                        R.fails.append(f"{bk.name(key)} sums a ratio/% row ({bk.name((sh, hit[0], c))} = "
                                       f"{fmt(bk.val((sh, hit[0], c)))}) into a total of levels")
    R.info.append(f"{n} total cells tested against the rows above them")
    if R.fails and len(R.fails) > MAX_LIST:
        R.fails = R.fails[:MAX_LIST] + [f"... and {len(R.fails) - MAX_LIST} more"]
    return R


def check_hardcodes(bk: Book):
    R = Result("hardcodes", "No numbers typed inside formulas on Calculations/Outputs")
    hits = []
    for k in sorted(bk.m.cells):
        if bk.roles.get(k[0]) not in ("calc", "output", "check"):
            continue
        f = bk.m.formula(k)
        if not f:
            continue
        try:
            toks = Tokenizer(f).items
        except Exception:
            continue
        lits = [t.value for t in toks if t.type == Token.OPERAND and t.subtype == Token.NUMBER
                and float(t.value) not in (0.0, 1.0)]
        if lits:
            hits.append(f"{bk.name(k)}: {f}  -> {', '.join(lits)}")
    if hits:
        R.fails.append(f"{len(hits)} formulas contain typed numbers (move each to Inputs with a source):")
        R.fails += hits[:MAX_LIST]
        if len(hits) > MAX_LIST:
            R.fails.append(f"... and {len(hits) - MAX_LIST} more")
    return R


def typed_inputs_on_model_sheets(bk: Book):
    """Labelled numeric constants on calc/output sheets that are not controls/axes/pasted grids."""
    ctl = bk.controls()
    grid_cells = set()
    for g in find_pasted_grids(bk):
        for r in range(g["r1"] - 1, g["r2"] + 1):
            for c in range(g["c1"] - 1, g["c2"] + 1):
                grid_cells.add((g["sheet"], r, c))
    tl = {(sh, t["row"], c) for sh, ts in bk.timelines.items() for t in ts for c in t["cols"]}
    out = []
    for k in sorted(bk.m.cells):
        if bk.roles.get(k[0]) in ("calc", "output", "check") and bk.numeric_constant(k) \
                and k not in ctl and k not in grid_cells and k not in tl:
            out.append(k)
    return out


def check_sources(bk: Book):
    R = Result("sources", "Every input has a source, date, A-F/1-6 grade and a low/base/high range")
    m = bk.m
    if not bk.registers:
        typed = typed_inputs_on_model_sheets(bk)
        R.fails.append("no Inputs register with Source and Grade columns")
        if typed:
            R.fails.append(f"{len(typed)} typed numbers on model sheets have no source or grade, e.g. " +
                           "; ".join(f"{bk.name(k)} = {fmt(bk.val(k))}" for k in typed[:MAX_LIST]))
        return R
    for reg in bk.registers:
        sh = reg["sheet"]
        rows = sorted({k[1] for k in m.cells if k[0] == sh and k[1] > reg["header"]})
        for r in rows:
            nums = [k for k in m.cells if k[0] == sh and k[1] == r and is_num(bk.val(k))
                    and k[2] not in (reg["id"], reg["grade"], reg["date"], reg["bench_low"], reg["bench_high"], reg["bench_med"])]
            if not nums:
                continue
            lab = bk.row_label(sh, r, upto=max(3, (reg["source"] or 3) - 1))
            where = f"{sh}!row {r}" + (f" ({lab})" if lab else "")
            src = bk.text((sh, r, reg["source"]))
            grade = bk.text((sh, r, reg["grade"])).upper().replace(" ", "")
            date = bk.val((sh, r, reg["date"])) if reg["date"] else "n/a"
            typ = bk.text((sh, r, reg["type"])).lower() if reg["type"] else ""
            problems = []
            if not src:
                problems.append("no source")
            if reg["date"] and (date is None or date == ""):
                problems.append("no date")
            if not re.fullmatch(r"[A-F][1-6]", grade):
                problems.append(f"grade '{grade or '-'}' is not A-F + 1-6")
            if reg["low"] and reg["high"] and reg["base"] and typ not in ("actual", "constant", "control", "conversion",
                                                                          "unit", "plan"):
                lo, ba, hi = (bk.val((sh, r, reg[c])) for c in ("low", "base", "high"))
                if not all(is_num(x) for x in (lo, ba, hi)):
                    problems.append("missing low/base/high")
                elif not (min(lo, hi) <= ba <= max(lo, hi)):
                    problems.append(f"base {fmt(ba)} outside its range {fmt(lo)}-{fmt(hi)}")
                elif lo == hi:
                    problems.append("single point, no range (actuary rule: every driver gets a range)")
            if problems:
                R.fails.append(f"{where}: " + ", ".join(problems))
    typed = typed_inputs_on_model_sheets(bk)
    if typed:
        R.fails.append(f"{len(typed)} numbers typed on calculation/output sheets bypass the register, e.g. " +
                       "; ".join(f"{bk.name(k)} = {fmt(bk.val(k))}" for k in typed[:MAX_LIST]))
    if len(R.fails) > MAX_LIST + 2:
        R.fails = R.fails[:MAX_LIST] + [f"... and {len(R.fails) - MAX_LIST} more"]
    return R


def ts_rows(bk: Book, sheets):
    """Yield (sheet, row, timeline, [(col, period, marker, value)]) for numeric time-series rows."""
    for sh in sheets:
        rows = sorted({k[1] for k in bk.m.cells if k[0] == sh})
        for r in rows:
            t = bk.timeline_for(sh, r)
            if not t:
                continue
            pts = []
            for c, p, mk in zip(t["cols"], t["periods"], t["markers"]):
                x = bk.val((sh, r, c))
                if is_num(x):
                    pts.append((c, p, mk, x))
            if len(pts) >= 2:
                yield sh, r, t, pts


def row_flagged(bk, sh, r):
    for k in bk.m.cells:
        if k[0] == sh and k[1] == r:
            t = bk.text(k).lower()
            if t.startswith("flag") or "flag:" in t or "justified" in t:
                return True
    return False


def row_mentions(bk, sh, r, word):
    return any(word in bk.text(k).lower() for k in bk.m.cells if k[0] == sh and k[1] == r)


def check_penetration(bk: Book, cap):
    R = Result("penetration", "Penetration/share <= 100% of the market; volumes bounded by a ceiling")
    for sh, r, t, pts in ts_rows(bk, bk.sheets("calc", "output")):
        lab = bk.row_label(sh, r)
        if LBL_PEN.search(lab) and not re.search(r"factor|uplift|log share|sum of log", lab, re.I):
            over = [(p, x) for c, p, mk, x in pts if x > 1 + 1e-9]
            if over and not all(x > 1 for _, _, _, x in pts):
                R.fails.append(f"{sh}!row {r} ({lab}) exceeds 100%: " +
                               ", ".join(f"{p:g}: {x:.1%}" for p, x in over[:4]))
            elif over and max(x for _, _, _, x in pts) > 100:
                R.fails.append(f"{sh}!row {r} ({lab}) exceeds 100 (% scale)")
        if LBL_VOLUME.search(lab) and not LBL_RATIO.search(lab) and not LBL_SKIP_GROWTH.search(lab) \
                and not LBL_REPORTED.search(lab) and not re.search(r"back-?test", lab, re.I):
            step = (t["periods"][1] - t["periods"][0]) if len(t["periods"]) > 1 else 1
            pcap = (1 + cap) ** step - 1 if step > 0 else cap
            # reported history is not a projection: when the timeline marks projections (2025E), test only those
            has_proj = any(mk in PROJ_MK for mk in t["markers"]) and any(mk not in PROJ_MK for mk in t["markers"])
            fast = [(p, b / a - 1) for (_, _, _, a), (c, p, mk, b) in zip(pts, pts[1:])
                    if a > 0 and b > 0 and b / a - 1 > pcap and (not has_proj or mk in PROJ_MK)]
            formulas = " ".join(bk.m.formula((sh, r, c)) or "" for c, *_ in pts).upper()
            bounded = "MIN(" in formulas
            if fast and not bounded and not row_flagged(bk, sh, r):
                g = sum(x for _, x in fast) / len(fast)
                R.fails.append(f"{sh}!row {r} ({lab}) compounds at {g:.0%}/period "
                               f"({fmt(pts[0][3])} -> {fmt(pts[-1][3])}) with no ceiling (no MIN against capacity/market)")
    return R


def check_growth(bk: Book, cap):
    R = Result("growth", f"Year-on-year growth of level rows <= {cap:.0%} unless flagged")
    for sh, r, t, pts in ts_rows(bk, bk.sheets("calc", "output")):
        lab = bk.row_label(sh, r)
        if not lab or LBL_RATIO.search(lab) or LBL_SKIP_GROWTH.search(lab) or LBL_PEN.search(lab):
            continue
        if not is_num(pts[0][3]) or pts[0][3] == 0:
            continue  # incremental rows start at zero
        proj = ("E", "F", "B", "P")
        # projections are marked (2025E) only when the timeline also has actual columns
        has_proj = any(mk in proj for mk in t["markers"]) and any(mk not in proj for mk in t["markers"])
        if has_proj and all(mk in proj for _, _, mk, _ in pts):
            continue  # no actual/base-year value: an incremental or projection-only row
        if not has_proj and pts[0][0] != t["cols"][0]:
            continue  # starts after the timeline begins: incremental row
        step = (t["periods"][1] - t["periods"][0]) if len(t["periods"]) > 1 else 1
        pcap = (1 + cap) ** step - 1 if step > 0 else cap
        flagged = row_flagged(bk, sh, r)
        for (c0, p0, m0, a), (c, p, mk, b) in zip(pts, pts[1:]):
            if has_proj and mk not in proj:
                continue
            if a > 0 and b > 0 and b / a - 1 > pcap + 1e-12:
                msg = f"{sh}!{addr(r, c)} ({lab}): {b / a - 1:+.1%} in {p:g}" + \
                      (f" (cap {pcap:.1%} per {step:g}-yr period)" if step != 1 else "")
                (R.warns if flagged else R.fails).append(msg + (" - flagged" if flagged else ""))
    if len(R.fails) > MAX_LIST:
        R.fails = R.fails[:MAX_LIST] + [f"... and {len(R.fails) - MAX_LIST} more"]
    return R


def check_timeline(bk: Book):
    R = Result("timeline", "One consistent timeline across sheets")
    segs = []
    for sh, ts in bk.timelines.items():
        if bk.roles[sh] == "skip":
            continue
        for t in ts:
            ps = t["periods"]
            tag = f"{sh}!row {t['row']}"
            if len(set(ps)) < len(ps):
                d = sorted({p for p in ps if ps.count(p) > 1})
                R.fails.append(f"{tag}: period repeated ({', '.join(f'{p:g}' for p in d)}) - check the header labels")
            steps = [round(b - a, 4) for a, b in zip(ps, ps[1:])]
            if any(s <= 0 for s in steps):
                R.fails.append(f"{tag}: periods not increasing")
            elif len(set(steps)) > 1:
                R.fails.append(f"{tag}: uneven step between periods ({', '.join(f'{p:g}' for p in ps)})")
            segs.append((sh, t))
    # same period must sit in the same column everywhere (FAST: one timeline)
    seen = {}
    for sh, t in segs:
        steps = [round(b - a, 4) for a, b in zip(t["periods"], t["periods"][1:])]
        gran = min(steps) if steps and min(steps) > 0 else None
        for c, p in zip(t["cols"], t["periods"]):
            seen.setdefault((gran, p), []).append((sh, t["row"], c))
    clashes = {}
    for (gran, p), locs in seen.items():
        cols = {c for _, _, c in locs}
        if len(cols) > 1:
            key = tuple(sorted({(sh, r) for sh, r, _ in locs}))
            clashes.setdefault(key, []).append((p, locs))
    for key, items in clashes.items():
        p, locs = items[0]
        R.fails.append(f"period {p:g} sits in different columns: " +
                       ", ".join(f"{sh}!{get_column_letter(c)}{r}" for sh, r, c in locs[:4]) +
                       f" ({len(items)} periods misaligned)")
    if not segs:
        R.warns.append("no timeline header found")
    return R


def check_unused(bk: Book):
    R = Result("unused", "Every input is used (decay/flatten rates are applied)")
    m = bk.m
    dep = m.dependents()
    ctl = bk.controls()
    cands = []
    for reg in bk.registers:
        sh = reg["sheet"]
        for r in sorted({k[1] for k in m.cells if k[0] == sh and k[1] > reg["header"]}):
            cols = [reg[c] for c in ("base",) if reg[c]] or \
                   [k[2] for k in m.cells if k[0] == sh and k[1] == r and is_num(bk.val(k))
                    and k[2] not in (reg["id"], reg["grade"], reg["date"], reg["bench_low"], reg["bench_high"], reg["bench_med"])]
            keys = [(sh, r, c) for c in cols if is_num(bk.val((sh, r, c)))]
            if keys:
                cands.append(("row", keys))
    regsheets = {reg["sheet"] for reg in bk.registers}
    for k in sorted(m.cells):
        if k[0] in regsheets or bk.roles.get(k[0]) == "skip" or k in ctl or not bk.numeric_constant(k):
            continue
        if bk.roles.get(k[0]) == "input":
            # historic data rows: whole row unused
            continue
        if bk.label(k):
            cands.append(("cell", [k]))
    # input-type sheets without register: rows of data
    for sh in bk.sheets("input"):
        if sh in regsheets:
            continue
        for r in sorted({k[1] for k in m.cells if k[0] == sh}):
            keys = [k for k in m.cells if k[0] == sh and k[1] == r and bk.numeric_constant(k)]
            if keys:
                cands.append(("data", keys))
    grid_cells = {(g["sheet"], r, c) for g in find_pasted_grids(bk)
                  for r in range(g["r1"] - 1, g["r2"] + 1) for c in range(g["c1"] - 1, g["c2"] + 1)}
    tl = {(sh, t["row"], c) for sh, ts in bk.timelines.items() for t in ts for c in t["cols"]}
    unused_data = []
    for kind, keys in cands:
        if any(k in grid_cells or k in tl for k in keys):
            continue
        if any(dep.get(k) for k in keys):
            continue
        if all(bk.val(k) == 0 for k in keys):
            continue  # typed zero placeholders
        k = keys[0]
        lab = bk.label(k) or bk.row_label(k[0], k[1])
        msg = f"{bk.name(k) if kind == 'cell' else k[0] + '!row ' + str(k[1]) + ' (' + lab + ')'} = " \
              f"{fmt(bk.val(k))} is defined but never used"
        if LBL_DECAY.search(lab):
            msg += " - the decay/flatten rate is never applied"
        if kind == "data":
            unused_data.append(msg)
        else:
            R.fails.append(msg)
    if unused_data:
        R.warns.append(f"{len(unused_data)} rows of reference data are not used by the model, e.g. " +
                       "; ".join(unused_data[:4]))
    if len(R.fails) > MAX_LIST:
        R.fails = R.fails[:MAX_LIST] + [f"... and {len(R.fails) - MAX_LIST} more"]
    return R


def output_cells(bk: Book):
    outs = [k for k in bk.m.cells if bk.roles.get(k[0]) == "output" and (bk.m.is_formula(k) or k in bk.m.dt_index)]
    if not outs:
        outs = [k for k in bk.m.cells if bk.roles.get(k[0]) == "calc" and bk.m.is_formula(k)]
    return outs


def check_trace(bk: Book):
    R = Result("trace", "Every output traces back to sourced input cells (forensic chain)")
    m = bk.m
    outs = output_cells(bk)
    prec = m.precedents()
    seen, stack = set(outs), list(outs)
    while stack:
        k = stack.pop()
        for p in prec.get(k, ()):
            if p not in seen:
                seen.add(p)
                stack.append(p)
    ctl = bk.controls() | {(sh, t["row"], c) for sh, ts in bk.timelines.items() for t in ts for c in t["cols"]}
    leaves = [k for k in seen if bk.numeric_constant(k)]
    bad = []
    for k in sorted(leaves):
        if k in ctl:
            continue
        if bk.register_row(k) and bk.sourced(k):
            continue
        bad.append(k)
    R.info.append(f"{len(outs)} output cells, {len(leaves)} typed numbers in their chains")
    if bad:
        R.fails.append(f"{len(bad)} typed numbers feed the outputs without a sourced, graded register row, e.g. " +
                       "; ".join(f"{bk.name(k)} = {fmt(bk.val(k))}" for k in bad[:MAX_LIST]))
    dangling = [k for k in seen if k not in m.cells and k not in m.dt_index and bk.roles.get(k[0]) != "skip"]
    refs_to_blank = [k for k in dangling if k[0] in m.sheetnames and k not in ctl]
    if refs_to_blank and len(refs_to_blank) < 50:
        hit = [k for k in refs_to_blank if any(k in prec.get(o, ()) for o in seen if m.is_formula(o)
                                               and len(prec.get(o, ())) <= 3)]
        if hit:
            R.warns.append(f"{len(hit)} formulas point at empty cells, e.g. " + ", ".join(bk.name(k) for k in hit[:5]))
    return R


def check_checks_sheet(bk: Book):
    R = Result("checks_sheet", "Model's own Checks sheet exists and is all clear")
    sheets = bk.sheets("check")
    if not sheets:
        R.fails.append("no Checks sheet (FAST: integrity checks live in the model)")
        return R
    for sh in sheets:
        bad = []
        n = 0
        for k in bk.m.cells:
            if k[0] != sh or not bk.m.is_formula(k):
                continue
            x = bk.val(k)
            lab = bk.label(k)
            if isinstance(x, bool):
                n += 1
                if not x:
                    bad.append(f"{bk.name(k)} = FALSE")
            elif isinstance(x, str) and x.upper() in ("OK", "ERROR", "FAIL", "CHECK"):
                n += 1
                if x.upper() != "OK":
                    bad.append(f"{bk.name(k)} = {x}")
            elif re.search(r"master|errors?\b|flags?\b", lab, re.I) and is_num(x):
                n += 1
                if x != 0:
                    bad.append(f"{bk.name(k)} = {fmt(x)}")
        if n == 0:
            R.fails.append(f"{sh}: no check results found (use TRUE/FALSE, OK/ERROR or an error count)")
        R.fails += bad
        R.info.append(f"{sh}: {n} checks read")
    return R


def check_outside_view(bk: Book):
    R = Result("outside_view", "Outside view: every uplift has a reference class (p25 / median / p75); above p75 "
                               "only with a written justification; no lump-sum savings")
    if not bk.registers:
        R.fails.append("no Inputs register with Source and Grade columns - reference class cannot be verified")
        return R
    m = bk.m
    n_up = 0
    for reg in bk.registers:
        if not reg["type"]:
            if reg["low"] or reg["high"] or reg["bench_high"]:   # an assumptions table, not a sheet of actuals
                R.fails.append(f"{reg['sheet']}: register has no Type column - uplifts cannot be identified")
            continue
        sh = reg["sheet"]
        for r in sorted({k[1] for k in m.cells if k[0] == sh and k[1] > reg["header"]}):
            typ = bk.text((sh, r, reg["type"])).lower()
            ba = bk.val((sh, r, reg["base"])) if reg["base"] else None
            if not is_num(ba):
                continue
            lab = bk.row_label(sh, r, upto=max(3, (reg["source"] or 3) - 1))
            where = f"{sh}!row {r}" + (f" ({lab})" if lab else "")
            if typ == "cost" and ba < 0:
                R.fails.append(f"{where}: lump-sum saving typed as a negative cost ({fmt(ba)}) - model it from its "
                               f"drivers (volume x price, e.g. a build initiative of kind 'saving') so every driver "
                               f"is a sourced input")
            if typ not in ("uplift", "effect"):
                continue
            n_up += 1
            g = lambda c: bk.val((sh, r, reg[c])) if reg[c] else None
            bl, bm, bh = g("bench_low"), g("bench_med"), g("bench_high")
            src = bk.text((sh, r, reg["bench_src"])) if reg["bench_src"] else ""
            just = bk.text((sh, r, reg["justif"])) if reg["justif"] else ""
            if not (is_num(bl) and is_num(bh)):
                R.fails.append(f"{where}: uplift has no reference class (p25 and p75 of comparable outcomes)")
                continue
            if not src:
                R.fails.append(f"{where}: reference class has no source")
            if not is_num(bm):
                R.warns.append(f"{where}: no reference-class median given (the base should sit near it)")
            if src and not re.search(r"\bn\s*[=~]\s*\d+|\b\d+\s+(cases|comparables|studies|programmes|programs|"
                                     r"launches|data points|applications)", src, re.I):
                R.warns.append(f"{where}: reference class size not stated (write 'n = 7'); with fewer than 5 cases "
                               f"grade the uplift C4 or worse")
            if ba > bh and not just:
                R.fails.append(f"{where}: base {fmt(ba)} is above the reference-class upper quartile {fmt(bh)} "
                               f"with no written justification")
            elif ba > bh:
                R.info.append(f"{where}: above p75, justified: {just[:80]}")
    R.info.append(f"{n_up} uplifts checked against their reference class")
    return R


LBL_REV = re.compile(r"revenue|turnover|sales|ventas|ingresos|receita|umsatz", re.I)
LBL_REV_SKIP = re.compile(r"incremental|uplift|target|with strateg|back-?test|gap|growth|%|margin|share|\bper\b|"
                          r"cagr|calibration|\bcost|model|percentile|median|cumulative|pv\b|npv", re.I)
LBL_BASELINE = re.compile(r"baseline|do[- ]nothing|status quo|base case|organic|business as usual|\bbau\b", re.I)
LBL_REPORTED = re.compile(r"reported|actual|historic|\bhist\b", re.I)
PROJ_MK = ("E", "F", "B", "P")


def _pct(xs, k):
    s = sorted(xs)
    h = (len(s) - 1) * k
    lo = int(h)
    return s[lo] + (h - lo) * (s[min(lo + 1, len(s) - 1)] - s[lo])


LBL_SHOCK = re.compile(r"shock[ _-]?year", re.I)
LBL_BRIDGE = re.compile(r"^\s*bridge to history\s*:\s*(.*)$", re.I)


def shock_years(bk: Book):
    """Years flagged 1 on a 'Shock year' row of an input sheet: left out of the band with the rebound after them."""
    out = set()
    for sh, r, t, pts in ts_rows(bk, bk.sheets("input")):
        if LBL_SHOCK.search(bk.row_label(sh, r)):
            out |= {p for c, p, mk, x in pts if x == 1}
    return out


def history_bridge(bk: Book):
    """(start, [(step name, value)], where) from the 'Bridge to history:' rows, or None."""
    start, steps, where = None, [], None
    for sh in bk.sheets("calc", "output"):
        for r in sorted({k[1] for k in bk.m.cells if k[0] == sh}):
            lab = bk.row_label(sh, r)
            mt = LBL_BRIDGE.match(lab or "")
            if not mt:
                continue
            nums = [bk.val(k) for k in sorted(bk.m.cells) if k[0] == sh and k[1] == r and k[2] > 2
                    and is_num(bk.val(k))]
            if not nums:
                continue
            if mt.group(1).lower().startswith("start"):
                start, steps, where = nums[0], [], f"{sh}!row {r}"
            elif start is not None:
                steps.append((mt.group(1).strip(), nums[0]))
        if start is not None:
            return start, steps, where
    return None


def check_history(bk: Book, tol=HIST_TOL):
    R = Result("history", "Base case anchored in reported history: 3+ years of revenue; do-nothing growth inside "
                          "the company's historical band, or bridged to it step by step")
    best = None
    for sh, r, t, pts in ts_rows(bk, bk.sheets("input", "calc", "output")):
        lab = bk.row_label(sh, r)
        if not lab or not LBL_REV.search(lab) or LBL_REV_SKIP.search(lab) or LBL_BASELINE.search(lab):
            continue
        marked = any(t["markers"])
        if bk.roles[sh] != "input" and not (marked or LBL_REPORTED.search(lab)):
            continue
        hist = [(p, x) for c, p, mk, x in pts if not marked or mk not in PROJ_MK]
        if bk.roles[sh] != "input" and marked and not LBL_REPORTED.search(lab) and len(hist) < len(pts):
            continue  # a projection row that merely starts in an actual year
        if best is None or len(hist) > len(best[2]):
            best = (sh, r, hist, lab)
    if not best or len(best[2]) < 3:
        R.fails.append("no reported revenue history of 3+ years found (a 'Revenue'/'Turnover' row on the Actuals "
                       "sheet): the base case cannot be compared with the company's record - pull 5 years "
                       "(Orbis/SABI) and put them in the model")
        return R
    sh, r, hist, lab = best
    shock = shock_years(bk)
    g = [(b / a) ** (1 / (p1 - p0)) - 1 for (p0, a), (p1, b) in zip(hist, hist[1:]) if a > 0 and b > 0 and p1 > p0
         and p0 not in shock and p1 not in shock]
    clean = [x for x in hist if x[0] not in shock] if shock else hist
    if len(clean) < 2:
        R.fails.append(f"shock years {', '.join(f'{y:g}' for y in sorted(shock))} leave fewer than 2 clean years")
        return R
    (p0, v0), (p1, v1) = clean[0], clean[-1]
    cagr = (v1 / v0) ** (1 / (p1 - p0)) - 1 if v0 > 0 and v1 > 0 and p1 > p0 else None
    if len(g) < 2:
        R.fails.append(f"{sh}!row {r} ({lab}): fewer than 2 year-on-year changes - cannot form a historical band")
        return R
    q1, q3 = _pct(g, HIST_Q[0]), _pct(g, HIST_Q[1])
    lo = min(q1, cagr if cagr is not None else q1) - tol
    hi = max(q3, cagr if cagr is not None else q3) + tol
    R.info.append(f"history: {sh}!row {r} ({lab}) {p0:g}-{p1:g}, {len(hist)} years; CAGR "
                  f"{'n/a' if cagr is None else f'{cagr:.1%}'}; YoY growth p25 {q1:.1%}, p75 {q3:.1%}" +
                  (f"; shock years {', '.join(f'{y:g}' for y in sorted(shock))} and the rebound after them left out"
                   if shock else ""))
    if len(hist) < 5:
        R.warns.append(f"only {len(hist)} years of reported revenue - 5 give a sturdier band (pull them from the "
                       f"IE library)")
    cands = []
    for sh2, r2, t, pts in ts_rows(bk, bk.sheets("calc", "output")):
        lab2 = bk.row_label(sh2, r2)
        if not (LBL_BASELINE.search(lab2) and LBL_REV.search(lab2)) or \
                re.search(r"incremental|growth|%|cagr|gap|calibration|uplift|band", lab2, re.I):
            continue
        marked = any(t["markers"])
        act = [x for x in pts if marked and x[2] not in PROJ_MK]
        proj = [x for x in pts if not marked or x[2] in PROJ_MK]
        start = act[-1] if act else pts[0]
        end = proj[-1] if proj else pts[-1]
        if end[1] <= start[1] or start[3] <= 0 or end[3] <= 0:
            continue
        cands.append((abs(end[3]), ":" not in lab2, sh2, r2, lab2, start, end))
    if not cands:
        R.fails.append("no baseline (do-nothing) revenue projection found (a 'Baseline revenue' row on Calculations "
                       "or Outputs) - the base case cannot be compared with history")
        return R
    _, _, sh2, r2, lab2, (_, ps, _, vs), (_, pe, _, ve) = max(cands)
    bg = (ve / vs) ** (1 / (pe - ps)) - 1
    where = f"{sh2}!row {r2} ({lab2})"
    msg = (f"{where} grows {bg:.1%}/yr {ps:g}-{pe:g}; the company's own record: CAGR "
           f"{'n/a' if cagr is None else f'{cagr:.1%}'}, YoY growth p25-p75 {q1:.1%} to {q3:.1%} (band {lo:.1%} to "
           f"{hi:.1%})")
    if lo - 1e-12 <= bg <= hi + 1e-12:
        R.info.append(msg + " - inside")
    else:
        br = history_bridge(bk)
        if br is None:
            R.fails.append(msg + " - outside: tie the do-nothing growth to history (e.g. add inflation if history is "
                                 "nominal), or bridge it: rows 'Bridge to history: start' (the historical CAGR or "
                                 "median) + one row per step (meta.baseline_bridge)" +
                           (". A written justification alone is not evidence" if row_flagged(bk, sh2, r2) or
                            row_mentions(bk, sh2, r2, "justif") else ""))
        else:
            start, steps, where_b = br
            med = _pct(g, 0.5)
            end = start + sum(v for _, v in steps)
            if not any(x is not None and abs(start - x) <= 1e-4 for x in (cagr, med)):
                R.fails.append(msg + f" - the bridge ({where_b}) starts at {start:.1%}, which is neither the "
                                     f"historical CAGR nor the median YoY growth")
            elif abs(end - bg) > tol + 1e-12:
                R.fails.append(msg + f" - the bridge ({where_b}) reaches {end:.1%} = {start:.1%} " +
                               " ".join(f"{v:+.1%} {n}" for n, v in steps) + f", {abs(end - bg):.1%} away from the "
                               f"baseline (tolerance {tol:.0%})")
            else:
                R.warns.append(msg + f" - outside, bridged ({where_b}): {start:.1%} " +
                               " ".join(f"{v:+.1%} ({n})" for n, v in steps) + f" = {end:.1%}")
    linked = 0
    for reg in bk.registers:
        if reg["base"]:
            for k in bk.m.cells:
                if k[0] == reg["sheet"] and k[2] == reg["base"] and k[1] > reg["header"]:
                    f = bk.m.formula(k) or ""
                    if f and not f.startswith("=" + reg["sheet"]):
                        linked += 1
    if bk.registers and not linked:
        R.warns.append("no register input takes its value from the actuals (base-year levels and growth typed by "
                       "hand) - link them so the base case updates with the history")
    else:
        R.info.append(f"{linked} register inputs linked to history")
    return R


def check_revision(path, version_a, spec_a, proposals):
    """Check 14 (revised models only): the revised workbook differs from version A only by the proposals the
    team accepted, and version A itself is unchanged."""
    import revisions as rv
    R = Result("revision", "Revised model = version A + accepted proposals only; version A untouched")
    try:
        sa = json.load(open(spec_a))
        props = json.load(open(proposals))
        acc = rv.accepted(props)
    except (OSError, ValueError, rv.ProposalError) as e:
        R.fails.append(f"cannot read revision inputs: {e}")
        return R
    if not acc:
        R.fails.append("no proposal has decision.status = 'accepted' (with 'by' and 'on'): a revised model cannot "
                       "exist without a recorded team decision")
        return R
    for p in acc:
        for prob in rv.validate(sa, p):
            R.fails.append(f"accepted proposal not valid: {prob}")
    sib = os.path.splitext(path)[0] + "_spec.json"
    letter = json.load(open(sib)).get("meta", {}).get("version", "B") if os.path.exists(sib) else "B"
    tmp = tempfile.mkdtemp(prefix="imrev_")
    exp_a, exp_b = os.path.join(tmp, "expected_A.xlsx"), os.path.join(tmp, "expected_B.xlsx")
    try:
        rv.build_quiet(sa, exp_a)
        rv.build_quiet(rv.apply_proposals(sa, acc, version=letter), exp_b)
    except Exception as e:  # noqa: BLE001 - any rebuild failure is a check failure
        R.fails.append(f"could not rebuild from spec A + accepted proposals: {e}")
        return R
    da = rv.diff_workbooks(version_a, exp_a, limit=8)
    if da:
        R.fails.append(f"version A ({os.path.basename(version_a)}) no longer matches its spec - edited after the "
                       f"build: " + "; ".join(f"{s_}!{c}" for s_, c, *_ in da))
    db = rv.diff_workbooks(path, exp_b, limit=12)
    if db:
        R.fails.append("changes not explained by the accepted proposals (hand edits or proposals the team did not "
                       "accept): " + "; ".join(f"{s_}!{c} is {str(x)[:40]!r}, expected {str(y)[:40]!r}"
                                               for s_, c, x, y in db))
    pend = [p["id"] for p in props["proposals"] if p not in acc]
    R.info.append(f"accepted: {', '.join(p['id'] for p in acc)}; not applied: {', '.join(pend) or 'none'}")
    R.info.append(f"{len(rv.diff_workbooks(version_a, path, limit=1000))} cells differ between version A and this one")
    return R


def trace_cell(bk: Book, ref, depth=6):
    sheet, a = split_sheet(ref)
    sh = bk.m.sheet(sheet) if sheet else bk.m.sheetnames[0]
    rc = parse_cell(a)
    if not sh or not rc:
        print(f"cannot parse {ref}")
        return
    prec = bk.m.precedents()
    printed = set()

    def show(k, d):
        f = bk.m.formula(k)
        reg = bk.register_row(k)
        extra = ""
        if reg and bk.numeric_constant(k):
            r = k[1]
            extra = (f"  <- source: {bk.text((k[0], r, reg['source']))} | grade "
                     f"{bk.text((k[0], r, reg['grade']))}")
            if reg["date"]:
                extra += f" | date {bk.val((k[0], r, reg['date']))}"
            if reg["link"]:
                extra += f" | {bk.text((k[0], r, reg['link']))}"
        print("  " * d + f"{bk.name(k)} = {fmt(bk.val(k))}" + (f"   {f}" if f else "") + extra)
        if k in printed or d >= depth:
            return
        printed.add(k)
        ps = sorted(prec.get(k, ()))
        if len(ps) > 12:
            print("  " * (d + 1) + f"... {len(ps)} precedents (range); showing typed inputs only")
            ps = [p for p in ps if bk.numeric_constant(p)][:12]
        for p in ps:
            show(p, d + 1)
    run_big_stack(show, (sh,) + rc, 0)


# ----------------------------------------------------------------------------- main
def run(path, engine="auto", cap=GROWTH_CAP, revision=None):
    vals = Values(path, engine)
    bk = Book(vals)
    results = [
        check_recalc(bk),
        check_sensitivity(bk),
        check_totals(bk),
        check_hardcodes(bk),
        check_sources(bk),
        check_penetration(bk, cap),
        check_growth(bk, cap),
        check_timeline(bk),
        check_unused(bk),
        check_trace(bk),
        check_checks_sheet(bk),
        check_outside_view(bk),
        check_history(bk),
    ]
    if revision:
        results.append(check_revision(path, *revision))
    return bk, results


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx")
    ap.add_argument("--engine", default="auto", choices=["auto", "libreoffice", "formulas", "builtin"])
    ap.add_argument("--growth-cap", type=float, default=GROWTH_CAP, help="max YoY growth of level rows (0.25 = 25%%)")
    ap.add_argument("--json", help="write the report as JSON to this path")
    ap.add_argument("--trace", help="print the source chain of one cell, e.g. 'Outputs!E8'")
    ap.add_argument("--verbose", action="store_true", help="also print info lines")
    ap.add_argument("--version-a", help="revised models: the version A workbook (adds check 14)")
    ap.add_argument("--spec-a", help="revised models: the spec version A was built from")
    ap.add_argument("--proposals", help="revised models: proposals.json with the team's decisions")
    a = ap.parse_args()
    if not os.path.exists(a.xlsx):
        sys.exit(f"file not found: {a.xlsx}")
    rev = None
    if a.version_a or a.spec_a or a.proposals:
        if not (a.version_a and a.spec_a and a.proposals):
            sys.exit("revision mode needs all three: --version-a, --spec-a, --proposals")
        rev = (a.version_a, a.spec_a, a.proposals)
    bk, results = run(a.xlsx, a.engine, a.growth_cap, rev)
    if a.trace:
        trace_cell(bk, a.trace)
        return
    print(f"impact-model checker: {os.path.basename(a.xlsx)}")
    print(f"recalculated with: {bk.v.engine}")
    for i, r in enumerate(results, 1):
        print(f"\n[{r.status}] {i:>2}. {r.title}")
        for f in r.fails:
            print(f"      x {f}")
        for w in r.warns:
            print(f"      ! {w}")
        if a.verbose:
            for x in r.info:
                print(f"      . {x}")
    npass = sum(r.status != "FAIL" for r in results)
    print(f"\nRESULT: {npass}/{len(results)} checks passed" + ("" if npass == len(results) else " - fix every x above"))
    if a.json:
        with open(a.json, "w") as f:
            json.dump({"file": a.xlsx, "engine": bk.v.engine, "checks": [r.as_dict() for r in results],
                       "passed": npass, "total": len(results)}, f, indent=2, default=str)
    sys.exit(0 if npass == len(results) else 1)


if __name__ == "__main__":
    main()
